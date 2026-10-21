import csv
import io
import zipfile
from datetime import date
from django.db import transaction
from openpyxl import load_workbook
from desk.models import Account, Audit, Product, ProductMapping, SalesDaily, Thread


class ImportProblem(ValueError):
    pass


def read_table(upload):
    data = upload.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise ImportProblem('Максимальный размер файла — 5 МБ.')
    if upload.name.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024:
                    raise ImportProblem('Распакованный Excel превышает 20 МБ.')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            sheet = book.active
            iterator = sheet.iter_rows(values_only=True)
            header = [str(x or '').strip() for x in next(iterator)]
            rows = []
            for values in iterator:
                if not any(v is not None for v in values): continue
                if any(isinstance(v, str) and v.startswith('=') for v in values):
                    raise ImportProblem('Формулы в Excel не поддерживаются; вставьте значения.')
                rows.append(dict(zip(header, values)))
                if len(rows) > 5000: raise ImportProblem('Не более 5000 строк за импорт.')
            book.close()
        except ImportProblem: raise
        except Exception as e: raise ImportProblem('Не удалось прочитать Excel.') from e
    elif upload.name.lower().endswith('.csv'):
        try:
            text = data.decode('utf-8-sig')
            delimiter = ';' if text.splitlines()[0].count(';') > text.splitlines()[0].count(',') else ','
            reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
            header = reader.fieldnames or []
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > 5000: raise ImportProblem('Не более 5000 строк за импорт.')
        except (UnicodeError, csv.Error, IndexError) as e:
            raise ImportProblem('Нужен CSV в UTF-8 с заголовком.') from e
    else:
        raise ImportProblem('Поддерживаются CSV и XLSX.')
    if len(header) != len(set(header)):
        raise ImportProblem('Повторяющиеся заголовки столбцов.')
    if not rows: raise ImportProblem('Файл не содержит данных.')
    return [{str(k).strip(): str(v if v is not None else '').strip() for k, v in row.items()} for row in rows]


def required(row, names, line):
    missing = [n for n in names if not row.get(n)]
    if missing: raise ImportProblem(f'Строка {line}: заполните {", ".join(missing)}.')


@transaction.atomic
def import_mapping(company, upload, actor):
    rows = read_table(upload)
    seen = set()
    for line, row in enumerate(rows, 2):
        required(row, ['ProductID', 'Product Name', 'Brand', 'Volume', 'Concentration', 'Platform', 'Account', 'SKU', 'External Product ID', 'Legal Entity'], line)
        account = Account.objects.filter(company=company, name=row['Account']).first()
        if not account: raise ImportProblem(f'Строка {line}: кабинет не принадлежит выбранной компании.')
        platform = {'wildberries': 'wb', 'wb': 'wb', 'ym': 'ym', 'яндекс маркет': 'ym'}.get(row['Platform'].lower())
        if platform != account.platform: raise ImportProblem(f'Строка {line}: площадка не совпадает с кабинетом.')
        if row['Legal Entity'] != company.legal_entity:
            raise ImportProblem(f'Строка {line}: юридическое лицо не совпадает.')
        key = (account.pk, row['SKU'])
        if key in seen: raise ImportProblem(f'Строка {line}: повтор кабинета и SKU.')
        seen.add(key)
        if any(len(row[k]) > n for k, n in [('ProductID', 100), ('Product Name', 200), ('Brand', 100), ('Volume', 40), ('Concentration', 40), ('SKU', 120), ('External Product ID', 120)]):
            raise ImportProblem(f'Строка {line}: слишком длинное значение.')
        values = {k: row[v] for k, v in [('name', 'Product Name'), ('brand', 'Brand'), ('volume', 'Volume'), ('concentration', 'Concentration')]}
        same_physical = Product.objects.filter(code=row['ProductID']).first()
        if same_physical and any(getattr(same_physical, k) != v for k, v in values.items()):
            raise ImportProblem(f'Строка {line}: ProductID уже относится к товару с другим описанием. Используйте единый каталог идентификаторов.')
        product, created = Product.objects.get_or_create(company=company, code=row['ProductID'], defaults=values)
        if not created and any(getattr(product, k) != v for k, v in values.items()):
            raise ImportProblem(f'Строка {line}: конфликт описания ProductID {product.code}.')
        mapping, created = ProductMapping.objects.get_or_create(account=account, sku=row['SKU'], defaults={'product': product, 'external_product_id': row['External Product ID']})
        if mapping.product_id != product.pk or mapping.external_product_id != row['External Product ID']:
            raise ImportProblem(f'Строка {line}: SKU уже связан с другим товаром или внешним ID.')
        Thread.objects.filter(account=account, sku=mapping.sku, product__isnull=True).update(product=product)
        if created:
            Audit.objects.create(company=company, actor=actor, action='mapping.import', object_id=str(mapping.pk), after=row)
    return len(rows)


@transaction.atomic
def import_sales(company, upload, actor):
    rows = read_table(upload)
    seen = set()
    for line, row in enumerate(rows, 2):
        required(row, ['Account', 'SKU', 'Date', 'Redemptions'], line)
        account = Account.objects.filter(company=company, name=row['Account']).first()
        if not account or not ProductMapping.objects.filter(account=account, sku=row['SKU']).exists():
            raise ImportProblem(f'Строка {line}: неизвестный кабинет / unmapped SKU.')
        try:
            day = date.fromisoformat(row['Date'][:10])
            if day > date.today(): raise ValueError()
            values = {k: int(row.get(v, '') or 0) for k, v in [('redemptions', 'Redemptions'), ('orders', 'Orders'), ('delivered', 'Delivered'), ('returns', 'Returns')]}
            if any(v < 0 for v in values.values()): raise ValueError()
        except ValueError as e:
            raise ImportProblem(f'Строка {line}: дата YYYY-MM-DD, количества — целые неотрицательные числа, дата не в будущем.') from e
        key = (account.pk, row['SKU'], day)
        if key in seen: raise ImportProblem(f'Строка {line}: повтор дневной строки продаж.')
        seen.add(key)
        old = SalesDaily.objects.filter(account=account, sku=row['SKU'], date=day).values(*values).first()
        item, _ = SalesDaily.objects.update_or_create(account=account, sku=row['SKU'], date=day, defaults=values)
        Audit.objects.create(company=company, actor=actor, action='sales.import', object_id=str(item.pk), before=old or {}, after=values)
    from desk.voc.baselines import refresh_baselines
    for product in Product.objects.filter(company=company): refresh_baselines(product, actor)
    return len(rows)
